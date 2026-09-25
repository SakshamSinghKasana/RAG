import docx
from pathlib import Path
from contextvault.core.models import ParsedDocument, DocumentSection
from contextvault.core.exceptions import ParseError
from contextvault.parsers.base import BaseParser

class DOCXParser(BaseParser):
    """Parser for DOCX files using python-docx."""
    
    def supported_extensions(self) -> set[str]:
        return {'.docx'}
        
    def parse(self, file_path: Path, file_id: str) -> ParsedDocument:
        try:
            doc = docx.Document(str(file_path))
        except Exception as e:
            raise ParseError(f"Failed to open DOCX {file_path}: {e}")
            
        title = doc.core_properties.title
        if not title:
            title = file_path.name
            
        sections = []
        current_title = "Document Start"
        current_content = []
        full_content = []
        
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
                
            full_content.append(text)
            
            if para.style.name.startswith('Heading'):
                if current_content:
                    sections.append(DocumentSection(
                        title=current_title,
                        content='\n'.join(current_content),
                        metadata={"heading": current_title}
                    ))
                current_title = text
                current_content = [text]
            else:
                current_content.append(text)
                
        if current_content:
            sections.append(DocumentSection(
                title=current_title,
                content='\n'.join(current_content),
                metadata={"heading": current_title}
            ))
            
        return ParsedDocument(
            file_id=file_id,
            title=title,
            content='\n\n'.join(full_content),
            sections=sections,
            metadata={"source_type": "docx"}
        )
