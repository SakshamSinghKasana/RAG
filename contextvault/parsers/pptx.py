from pptx import Presentation
from pathlib import Path
from contextvault.core.models import ParsedDocument, DocumentSection
from contextvault.core.exceptions import ParseError
from contextvault.parsers.base import BaseParser

class PPTXParser(BaseParser):
    """Parser for PPTX files using python-pptx."""
    
    def supported_extensions(self) -> set[str]:
        return {'.pptx'}
        
    def parse(self, file_path: Path, file_id: str) -> ParsedDocument:
        try:
            prs = Presentation(str(file_path))
        except Exception as e:
            raise ParseError(f"Failed to open PPTX {file_path}: {e}")
            
        title = prs.core_properties.title
        if not title:
            title = file_path.name
            
        sections = []
        full_content = []
        
        for i, slide in enumerate(prs.slides):
            slide_text = []
            
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_text.append(shape.text.strip())
                elif shape.has_table:
                    for row in shape.table.rows:
                        row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if row_text:
                            slide_text.append(" | ".join(row_text))
                            
            if slide_text:
                slide_content = '\n'.join(slide_text)
                full_content.append(slide_content)
                sections.append(DocumentSection(
                    title=f"Slide {i + 1}",
                    content=slide_content,
                    metadata={"slide": i + 1}
                ))
                
        return ParsedDocument(
            file_id=file_id,
            title=title,
            content='\n\n'.join(full_content),
            sections=sections,
            metadata={"source_type": "pptx"}
        )
