"""Data analysis and chart generation tool for CSV and Excel files in Context Vault."""

import csv
import logging
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from contextvault.core.vault import Vault

logger = logging.getLogger(__name__)


class ChartGenerator:
    """Extracts tabular data from vault files and generates professional visual charts."""

    @staticmethod
    def inspect_dataset(vault: Vault, relative_path: str) -> Dict[str, Any]:
        """Inspect a CSV or Excel dataset to determine columns, types, and summary statistics."""
        target_path = vault.absolute_path(relative_path)
        if not target_path.exists():
            raise FileNotFoundError(f"File not found in vault: {relative_path}")

        ext = target_path.suffix.lower()
        rows: List[List[Any]] = []
        headers: List[str] = []

        if ext in (".csv", ".tsv"):
            delimiter = "\t" if ext == ".tsv" else ","
            for enc in ("utf-8", "latin-1"):
                try:
                    with open(target_path, "r", encoding=enc, errors="replace") as f:
                        reader = csv.reader(f, delimiter=delimiter)
                        for row in reader:
                            if row and any(row):
                                rows.append([c.strip() for c in row])
                    break
                except Exception:
                    continue
            if rows:
                headers = rows[0]
                rows = rows[1:]
        elif ext == ".xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(target_path, read_only=True, data_only=True)
            sheet = wb.active
            if sheet:
                for i, row in enumerate(sheet.iter_rows(values_only=True)):
                    clean_row = [v for v in row]
                    if not any(clean_row):
                        continue
                    if i == 0:
                        headers = [str(c).strip() for c in clean_row if c is not None]
                    else:
                        rows.append(clean_row)
            wb.close()
        else:
            raise ValueError(f"Unsupported dataset format '{ext}'. Supported: .csv, .tsv, .xlsx")

        if not headers:
            return {"error": "No headers or data found in file", "columns": [], "row_count": 0}

                                   
        col_types: Dict[str, str] = {}
        summary_stats: Dict[str, Dict[str, float]] = {}
        samples: Dict[str, List[Any]] = {}

        for col_idx, col_name in enumerate(headers):
            values: List[Any] = []
            numeric_vals: List[float] = []

            for r in rows:
                if col_idx < len(r) and r[col_idx] is not None:
                    val = r[col_idx]
                    values.append(val)
                                             
                    try:
                        clean_num = str(val).replace(",", "").replace("$", "").replace("%", "").strip()
                        num = float(clean_num)
                        numeric_vals.append(num)
                    except (ValueError, TypeError):
                        pass

            samples[col_name] = values[:3]

            if len(numeric_vals) >= len(values) * 0.7 and len(numeric_vals) > 0:
                col_types[col_name] = "numeric"
                summary_stats[col_name] = {
                    "count": float(len(numeric_vals)),
                    "min": float(min(numeric_vals)),
                    "max": float(max(numeric_vals)),
                    "mean": float(sum(numeric_vals) / len(numeric_vals)),
                }
            else:
                col_types[col_name] = "categorical"

        return {
            "relative_path": relative_path,
            "filename": target_path.name,
            "row_count": len(rows),
            "columns": headers,
            "column_types": col_types,
            "summary_stats": summary_stats,
            "samples": samples,
        }

    @staticmethod
    def generate_chart(
        vault: Vault,
        relative_path: str,
        chart_type: str = "bar",
        x_column: Optional[str] = None,
        y_column: Optional[str] = None,
        title: Optional[str] = None,
        output_filename: Optional[str] = None,
    ) -> Dict[str, Any]:
        """Generate a high-quality chart PNG using matplotlib and save into Generated/charts/."""
        import matplotlib
        matplotlib.use("Agg")                           
        import matplotlib.pyplot as plt

        inspection = ChartGenerator.inspect_dataset(vault, relative_path)
        if "error" in inspection:
            raise ValueError(inspection["error"])

        columns = inspection["columns"]
        col_types = inspection["column_types"]

                                                               
        if not x_column:
                                             
            cats = [c for c in columns if col_types.get(c) == "categorical"]
            x_column = cats[0] if cats else columns[0]

        if not y_column:
                                         
            nums = [c for c in columns if col_types.get(c) == "numeric"]
            y_column = nums[0] if nums else (columns[1] if len(columns) > 1 else columns[0])

                          
        target_path = vault.absolute_path(relative_path)
        ext = target_path.suffix.lower()
        x_vals: List[Any] = []
        y_vals: List[float] = []

        if ext in (".csv", ".tsv"):
            delimiter = "\t" if ext == ".tsv" else ","
            with open(target_path, "r", encoding="utf-8", errors="replace") as f:
                reader = csv.DictReader(f, delimiter=delimiter)
                for row in reader:
                    raw_x = row.get(x_column)
                    raw_y = row.get(y_column)
                    if raw_x is not None and raw_y is not None:
                        try:
                            clean_y = str(raw_y).replace(",", "").replace("$", "").replace("%", "").strip()
                            y_vals.append(float(clean_y))
                            x_vals.append(str(raw_x).strip())
                        except (ValueError, TypeError):
                            continue
        elif ext == ".xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(target_path, read_only=True, data_only=True)
            sheet = wb.active
            headers: List[str] = []
            if sheet:
                for i, row in enumerate(sheet.iter_rows(values_only=True)):
                    if i == 0:
                        headers = [str(c).strip() for c in row if c is not None]
                    else:
                        row_dict = {headers[idx]: val for idx, val in enumerate(row) if idx < len(headers)}
                        raw_x = row_dict.get(x_column)
                        raw_y = row_dict.get(y_column)
                        if raw_x is not None and raw_y is not None:
                            try:
                                clean_y = str(raw_y).replace(",", "").replace("$", "").replace("%", "").strip()
                                y_vals.append(float(clean_y))
                                x_vals.append(str(raw_x).strip())
                            except (ValueError, TypeError):
                                continue
            wb.close()

        if not y_vals:
            raise ValueError(f"No valid numeric data could be extracted for Y column '{y_column}'.")

                                                                        
        if len(x_vals) > 25:
            x_vals = x_vals[:25]
            y_vals = y_vals[:25]

                                        
        charts_dir = vault.generated_dir / "charts"
        charts_dir.mkdir(parents=True, exist_ok=True)

        chart_name = output_filename or f"chart_{x_column}_{y_column}_{chart_type}.png"
        chart_name = "".join(c if c.isalnum() or c in "._-" else "_" for c in chart_name)
        if not chart_name.endswith(".png"):
            chart_name += ".png"

        output_path = charts_dir / chart_name

                                             
        fig, ax = plt.subplots(figsize=(9, 5.5), dpi=300)

                        
        primary_color = "#0284c7"                 
        accent_color = "#06b6d4"         
        bg_color = "#ffffff"
        grid_color = "#e2e8f0"

        fig.patch.set_facecolor(bg_color)
        ax.set_facecolor(bg_color)

        chart_type_lower = chart_type.lower()
        final_title = title or f"{y_column} by {x_column}"

        if chart_type_lower in ("bar", "column"):
            bars = ax.bar(x_vals, y_vals, color=primary_color, edgecolor="#0369a1", width=0.6, zorder=3)
                                                                   
            if len(bars) <= 15:
                for b in bars:
                    height = b.get_height()
                    ax.annotate(
                        f"{height:,.1f}" if height % 1 != 0 else f"{int(height):,}",
                        xy=(b.get_x() + b.get_width() / 2, height),
                        xytext=(0, 4),
                        textcoords="offset points",
                        ha="center", va="bottom", fontsize=8, color="#334155"
                    )
        elif chart_type_lower in ("line", "trend"):
            ax.plot(x_vals, y_vals, color=primary_color, marker="o", linewidth=2.5, markersize=6, zorder=3)
            ax.fill_between(range(len(x_vals)), y_vals, color=primary_color, alpha=0.12, zorder=2)
        elif chart_type_lower in ("scatter", "points"):
            ax.scatter(range(len(x_vals)), y_vals, color=primary_color, s=50, alpha=0.8, edgecolors="#0369a1", zorder=3)
            ax.set_xticks(range(len(x_vals)))
            ax.set_xticklabels(x_vals)
        elif chart_type_lower in ("pie", "donut"):
                                                 
            if len(x_vals) > 7:
                pie_x = x_vals[:6] + ["Other"]
                pie_y = y_vals[:6] + [sum(y_vals[6:])]
            else:
                pie_x, pie_y = x_vals, y_vals
            ax.pie(pie_y, labels=pie_x, autopct="%1.1f%%", startangle=140, colors=["#0284c7", "#0ea5e9", "#38bdf8", "#06b6d4", "#10b981", "#6366f1", "#94a3b8"])
        else:
                            
            ax.bar(x_vals, y_vals, color=primary_color, edgecolor="#0369a1", width=0.6, zorder=3)

        if chart_type_lower not in ("pie", "donut"):
            ax.set_title(final_title, fontsize=14, fontweight="bold", pad=15, color="#0f172a")
            ax.set_xlabel(x_column, fontsize=10, fontweight="bold", color="#334155", labelpad=8)
            ax.set_ylabel(y_column, fontsize=10, fontweight="bold", color="#334155", labelpad=8)
            ax.grid(True, linestyle="--", alpha=0.6, color=grid_color, zorder=1)
            ax.tick_params(axis="x", rotation=35 if len(x_vals) > 5 else 0, labelsize=9, colors="#334155")
            ax.tick_params(axis="y", labelsize=9, colors="#334155")
            ax.spines["top"].set_visible(False)
            ax.spines["right"].set_visible(False)
            ax.spines["left"].set_color("#cbd5e1")
            ax.spines["bottom"].set_color("#cbd5e1")
        else:
            ax.set_title(final_title, fontsize=14, fontweight="bold", pad=15, color="#0f172a")

        plt.tight_layout()
        plt.savefig(str(output_path), dpi=300, facecolor=bg_color, bbox_inches="tight")
        plt.close(fig)

        rel_out = vault.relative_path(output_path)
        logger.info(f"Chart generated at: {rel_out}")

        return {
            "chart_type": chart_type,
            "title": final_title,
            "image_relative_path": rel_out,
            "image_absolute_path": str(output_path),
            "x_column": x_column,
            "y_column": y_column,
            "points_count": len(y_vals),
        }
